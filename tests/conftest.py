"""Reusable synthetic uploads for offline tests."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from core.ingestion import ingest_upload
from core.profiler import profile_dataset

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def valid_csv_bytes() -> bytes:
    """Return the small valid CSV fixture."""
    return (FIXTURES / "valid.csv").read_bytes()


@pytest.fixture
def missing_duplicates_csv_bytes() -> bytes:
    """Return the fixture containing nulls and a repeated row."""
    return (FIXTURES / "missing_duplicates.csv").read_bytes()


@pytest.fixture
def mixed_ambiguous_csv_bytes() -> bytes:
    """Return a fixture with mixed values and ambiguous date strings."""
    return (FIXTURES / "mixed_ambiguous.csv").read_bytes()


@pytest.fixture
def malformed_csv_bytes() -> bytes:
    """Return a CSV with an unterminated quoted field."""
    return (FIXTURES / "malformed.csv").read_bytes()


@pytest.fixture
def duplicate_columns_csv_bytes() -> bytes:
    """Return a CSV with duplicate and colliding headers."""
    return (FIXTURES / "duplicate_columns.csv").read_bytes()


@pytest.fixture
def over_row_limit_csv_bytes() -> bytes:
    """Return a CSV used to verify that row limits reject, not truncate."""
    return (FIXTURES / "over_row_limit.csv").read_bytes()


@pytest.fixture
def unsupported_file_bytes() -> bytes:
    """Return content using an unsupported JSON file type."""
    return (FIXTURES / "unsupported.json").read_bytes()


@pytest.fixture
def blank_csv_bytes() -> bytes:
    """Return a whitespace-only CSV fixture."""
    return (FIXTURES / "blank.csv").read_bytes()


@pytest.fixture
def multi_sheet_xlsx_bytes() -> bytes:
    """Build a multiple-sheet workbook entirely in memory."""
    workbook = Workbook()
    sales = workbook.active
    sales.title = "Sales"
    sales.append(["region", "amount", "formula", "mixed"])
    sales.append(["North", 10, "=1+1", 10])
    sales.append(["South", 20, "=1+1", "unknown"])

    inventory = workbook.create_sheet("Inventory")
    inventory.append(["enabled", "created"])
    inventory.append([True, date(2025, 1, 2)])

    payload = BytesIO()
    workbook.save(payload)
    workbook.close()
    return payload.getvalue()


@pytest.fixture
def analysis_dataset(valid_csv_bytes: bytes):
    """Return the deterministic sales table used by plan pipeline tests."""
    return ingest_upload("sales.csv", valid_csv_bytes)


@pytest.fixture
def analysis_column_ids(analysis_dataset):
    """Map source display names to stable table-local column identifiers."""
    table = profile_dataset(analysis_dataset).tables[0]
    return {column.name: column.column_id for column in table.columns}
