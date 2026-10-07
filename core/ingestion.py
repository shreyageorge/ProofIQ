"""Bounded, in-memory ingestion of CSV and XLSX uploads."""

from __future__ import annotations

import hashlib
import io
import zipfile
import zlib
from pathlib import PureWindowsPath
from typing import Sequence
from xml.etree.ElementTree import ParseError

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from core.schema import stable_id, validate_column_names
from models.schemas import (
    CellValue,
    DatasetFormat,
    DatasetManifest,
    IngestionError,
    IngestionErrorCode,
    IngestionLimits,
    IngestedDataset,
    TableRef,
    TableSnapshot,
)

DEFAULT_INGESTION_LIMITS = IngestionLimits()
_XLSX_PARSER_ERRORS = (
    InvalidFileException,
    IndexError,
    KeyError,
    NotImplementedError,
    OSError,
    ParseError,
    RuntimeError,
    TypeError,
    ValueError,
    zipfile.BadZipFile,
    EOFError,
    zlib.error,
)


def ingest_upload(
    filename: str,
    content: bytes,
    *,
    limits: IngestionLimits = DEFAULT_INGESTION_LIMITS,
    selected_sheets: Sequence[str] | None = None,
) -> IngestedDataset:
    """Validate and ingest one CSV/XLSX byte payload without writing it to disk."""
    if not isinstance(content, bytes) or not content:
        raise IngestionError(
            IngestionErrorCode.INVALID_FILE, "The uploaded file is empty or invalid."
        )
    display_name = _safe_display_name(filename)
    suffix = PureWindowsPath(display_name).suffix.casefold()
    if suffix == ".csv":
        file_format = DatasetFormat.CSV
    elif suffix == ".xlsx":
        file_format = DatasetFormat.XLSX
    else:
        raise IngestionError(
            IngestionErrorCode.UNSUPPORTED_FORMAT,
            "Only CSV (.csv) and Excel workbooks (.xlsx) are supported.",
        )
    if len(content) > limits.max_file_size_bytes:
        raise IngestionError(
            IngestionErrorCode.FILE_TOO_LARGE,
            f"The file exceeds the {limits.max_file_size_bytes}-byte upload limit.",
        )

    content_sha256 = hashlib.sha256(content).hexdigest()
    dataset_id = f"ds_{content_sha256[:20]}"
    source_file_id = f"file_{content_sha256[:20]}"

    if file_format == DatasetFormat.CSV:
        if selected_sheets is not None:
            raise IngestionError(
                IngestionErrorCode.INVALID_FILE,
                "Sheet selection is only supported for XLSX workbooks.",
            )
        snapshots, refs = _read_csv(
            content, dataset_id, source_file_id, limits
        )
    else:
        snapshots, refs = _read_xlsx(
            content,
            dataset_id,
            source_file_id,
            limits,
            selected_sheets,
        )

    manifest = DatasetManifest(
        dataset_id=dataset_id,
        content_sha256=content_sha256,
        original_name=display_name,
        file_format=file_format,
        tables=refs,
        limits_applied=limits,
    )
    return IngestedDataset(manifest=manifest, tables=snapshots)


def _read_csv(
    content: bytes,
    dataset_id: str,
    source_file_id: str,
    limits: IngestionLimits,
) -> tuple[tuple[TableSnapshot, ...], tuple[TableRef, ...]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise IngestionError(
            IngestionErrorCode.INVALID_FILE,
            "The CSV must use UTF-8 or UTF-8 with a byte-order mark.",
        ) from None

    rows: list[tuple[CellValue, ...]] = []
    row_ids: list[str] = []
    source_row_numbers: list[int] = []
    columns: tuple[str, ...] | None = None
    source_row_number = 0
    table_id = stable_id("tbl", f"{dataset_id}\0csv")

    try:
        with pd.read_csv(
            io.StringIO(text),
            header=None,
            dtype=str,
            keep_default_na=False,
            na_filter=False,
            skip_blank_lines=False,
            engine="python",
            on_bad_lines="error",
            chunksize=min(limits.max_rows + 1, 10_000),
        ) as reader:
            for chunk in reader:
                if columns is None:
                    if chunk.empty:
                        continue
                    _check_column_limit(chunk.shape[1], limits)
                    raw_headers = chunk.iloc[0].tolist()
                    headers = [
                        None if pd.isna(value) or value == "" else value
                        for value in raw_headers
                    ]
                    columns = _validated_headers(headers)
                    _check_column_limit(len(columns), limits)
                    source_row_number += 1
                    data = chunk.iloc[1:]
                else:
                    data = chunk

                for raw_row in data.itertuples(index=False, name=None):
                    source_row_number += 1
                    if all(pd.isna(value) or value == "" for value in raw_row):
                        continue
                    values = tuple(
                        None if pd.isna(value) or value == "" else str(value)
                        for value in raw_row
                    )
                    rows.append(values)
                    row_ids.append(_row_id(table_id, source_row_number))
                    source_row_numbers.append(source_row_number)
                    _check_table_size(len(rows), len(columns), limits)
    except IngestionError:
        raise
    except (pd.errors.EmptyDataError, pd.errors.ParserError, ValueError):
        raise IngestionError(
            IngestionErrorCode.INVALID_FILE, "The CSV structure is malformed."
        ) from None
    if columns is None:
        raise IngestionError(
            IngestionErrorCode.INVALID_FILE, "The CSV does not contain a header row."
        )
    if not rows:
        raise IngestionError(
            IngestionErrorCode.EMPTY_DATASET, "The CSV does not contain any data rows."
        )

    snapshot, ref = _make_table(
        table_id=table_id,
        display_name="CSV",
        source_file_id=source_file_id,
        source_sheet=None,
        columns=columns,
        rows=rows,
        row_ids=row_ids,
        source_row_numbers=source_row_numbers,
    )
    return (snapshot,), (ref,)


def _read_xlsx(
    content: bytes,
    dataset_id: str,
    source_file_id: str,
    limits: IngestionLimits,
    selected_sheets: Sequence[str] | None,
) -> tuple[tuple[TableSnapshot, ...], tuple[TableRef, ...]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if (
                sum(info.file_size for info in infos)
                > limits.max_xlsx_uncompressed_bytes
            ):
                raise IngestionError(
                    IngestionErrorCode.LIMIT_EXCEEDED,
                    "The workbook exceeds the uncompressed size limit.",
                )
            if any(info.flag_bits & 0x1 for info in infos):
                raise IngestionError(
                    IngestionErrorCode.INVALID_FILE,
                    "Encrypted XLSX workbooks are not supported.",
                )
            required_parts = {"[Content_Types].xml", "xl/workbook.xml"}
            if not required_parts.issubset({info.filename for info in infos}):
                raise IngestionError(
                    IngestionErrorCode.INVALID_FILE,
                    "The uploaded file is not a valid XLSX workbook.",
                )
            if archive.testzip() is not None:
                raise IngestionError(
                    IngestionErrorCode.INVALID_FILE,
                    "The XLSX workbook contains a damaged archive entry.",
                )
    except zipfile.BadZipFile:
        raise IngestionError(
            IngestionErrorCode.INVALID_FILE, "The XLSX workbook is malformed."
        ) from None
    except (EOFError, NotImplementedError, RuntimeError, zlib.error):
        raise IngestionError(
            IngestionErrorCode.INVALID_FILE, "The XLSX workbook is malformed."
        ) from None

    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except _XLSX_PARSER_ERRORS:
        raise IngestionError(
            IngestionErrorCode.INVALID_FILE, "The XLSX workbook is malformed."
        ) from None

    try:
        sheet_names = tuple(workbook.sheetnames)
        if len(sheet_names) > limits.max_xlsx_sheets:
            raise IngestionError(
                IngestionErrorCode.LIMIT_EXCEEDED,
                f"The workbook exceeds the {limits.max_xlsx_sheets}-sheet limit.",
            )
        if selected_sheets is None:
            chosen_names = sheet_names
        else:
            chosen_names = tuple(selected_sheets)
            if len(set(chosen_names)) != len(chosen_names):
                raise IngestionError(
                    IngestionErrorCode.INVALID_FILE,
                    "A worksheet may only be selected once.",
                )
            unknown = tuple(name for name in chosen_names if name not in sheet_names)
            if unknown:
                raise IngestionError(
                    IngestionErrorCode.SHEET_NOT_FOUND,
                    f"Worksheet not found: {unknown[0]!r}.",
                )
        if not chosen_names:
            raise IngestionError(
                IngestionErrorCode.EMPTY_DATASET,
                "Select at least one worksheet to ingest.",
            )

        snapshots: list[TableSnapshot] = []
        refs: list[TableRef] = []
        for sheet_name in chosen_names:
            worksheet = workbook[sheet_name]
            if worksheet.max_column > limits.max_columns:
                raise IngestionError(
                    IngestionErrorCode.LIMIT_EXCEEDED,
                    f"Worksheet {sheet_name!r} exceeds the column limit.",
                )
            if max(worksheet.max_row - 1, 0) > limits.max_rows:
                raise IngestionError(
                    IngestionErrorCode.LIMIT_EXCEEDED,
                    f"Worksheet {sheet_name!r} exceeds the row limit.",
                )
            table_id = stable_id(
                "tbl",
                f"{dataset_id}\0xlsx\0{sheet_names.index(sheet_name)}\0{sheet_name}",
            )
            snapshot, ref = _read_worksheet(
                worksheet,
                table_id,
                source_file_id,
                sheet_name,
                limits,
            )
            snapshots.append(snapshot)
            refs.append(ref)
        return tuple(snapshots), tuple(refs)
    except IngestionError:
        raise
    except _XLSX_PARSER_ERRORS:
        raise IngestionError(
            IngestionErrorCode.INVALID_FILE, "The XLSX workbook is malformed."
        ) from None
    finally:
        workbook.close()


def _read_worksheet(
    worksheet: object,
    table_id: str,
    source_file_id: str,
    sheet_name: str,
    limits: IngestionLimits,
) -> tuple[TableSnapshot, TableRef]:
    try:
        iterator = worksheet.iter_rows(values_only=True)  # type: ignore[attr-defined]
        raw_headers = next(iterator)
    except StopIteration:
        raise IngestionError(
            IngestionErrorCode.EMPTY_DATASET,
            f"Worksheet {sheet_name!r} does not contain a header row.",
        ) from None

    trimmed_headers = list(raw_headers)
    while trimmed_headers and trimmed_headers[-1] is None:
        trimmed_headers.pop()
    if not trimmed_headers:
        raise IngestionError(
            IngestionErrorCode.EMPTY_DATASET,
            f"Worksheet {sheet_name!r} does not contain a header row.",
        )
    columns = _validated_headers(trimmed_headers)
    _check_column_limit(len(columns), limits)

    rows: list[tuple[CellValue, ...]] = []
    row_ids: list[str] = []
    source_row_numbers: list[int] = []
    for source_row_number, raw_row in enumerate(iterator, start=2):
        values = list(raw_row)
        while values and values[-1] is None:
            values.pop()
        if not values:
            continue
        if len(values) > len(columns):
            raise IngestionError(
                IngestionErrorCode.INVALID_SCHEMA,
                f"Worksheet {sheet_name!r}, row {source_row_number} has values "
                "beyond its header.",
            )
        padded = values + [None] * (len(columns) - len(values))
        if all(value is None for value in padded):
            continue
        rows.append(tuple(padded))
        row_ids.append(_row_id(table_id, source_row_number))
        source_row_numbers.append(source_row_number)
        _check_table_size(len(rows), len(columns), limits)

    if not rows:
        raise IngestionError(
            IngestionErrorCode.EMPTY_DATASET,
            f"Worksheet {sheet_name!r} does not contain any data rows.",
        )

    return _make_table(
        table_id=table_id,
        display_name=sheet_name,
        source_file_id=source_file_id,
        source_sheet=sheet_name,
        columns=columns,
        rows=rows,
        row_ids=row_ids,
        source_row_numbers=source_row_numbers,
    )


def _validated_headers(values: Sequence[object]) -> tuple[str, ...]:
    try:
        return validate_column_names(values)
    except ValueError as exc:
        raise IngestionError(
            IngestionErrorCode.INVALID_SCHEMA, str(exc)
        ) from None


def _check_column_limit(column_count: int, limits: IngestionLimits) -> None:
    if column_count > limits.max_columns:
        raise IngestionError(
            IngestionErrorCode.LIMIT_EXCEEDED,
            f"The table exceeds the {limits.max_columns}-column limit.",
        )


def _check_table_size(
    row_count: int, column_count: int, limits: IngestionLimits
) -> None:
    if row_count > limits.max_rows:
        raise IngestionError(
            IngestionErrorCode.LIMIT_EXCEEDED,
            f"The table exceeds the {limits.max_rows}-row limit.",
        )
    if row_count * column_count > limits.max_cells:
        raise IngestionError(
            IngestionErrorCode.LIMIT_EXCEEDED,
            f"The table exceeds the {limits.max_cells}-cell limit.",
        )


def _make_table(
    *,
    table_id: str,
    display_name: str,
    source_file_id: str,
    source_sheet: str | None,
    columns: tuple[str, ...],
    rows: Sequence[tuple[CellValue, ...]],
    row_ids: Sequence[str],
    source_row_numbers: Sequence[int],
) -> tuple[TableSnapshot, TableRef]:
    snapshot = TableSnapshot(
        table_id=table_id,
        columns=columns,
        rows=tuple(rows),
        row_ids=tuple(row_ids),
        source_row_numbers=tuple(source_row_numbers),
    )
    ref = TableRef(
        table_id=table_id,
        display_name=display_name,
        row_count=len(snapshot.rows),
        column_count=len(columns),
        source_file_id=source_file_id,
        source_sheet=source_sheet,
    )
    return snapshot, ref


def _row_id(table_id: str, source_row_number: int) -> str:
    return stable_id("row", f"{table_id}\0{source_row_number}")


def _safe_display_name(filename: str) -> str:
    if not isinstance(filename, str) or not filename.strip():
        raise IngestionError(
            IngestionErrorCode.INVALID_FILE, "A non-empty filename is required."
        )
    name = PureWindowsPath(filename).name
    safe_name = "".join(character for character in name if character.isprintable())
    if not safe_name or safe_name in {".", ".."}:
        raise IngestionError(
            IngestionErrorCode.INVALID_FILE, "The uploaded filename is invalid."
        )
    return safe_name
