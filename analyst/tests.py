from pathlib import Path
from io import BytesIO
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook

from .services import create_plan, execute_plan, load_tables, profile_tables, validate_code


class AnalystCoreTests(TestCase):
    def setUp(self):
        sample = Path(__file__).resolve().parent.parent / "sample_sales.csv"
        self.tables, self.sources = load_tables([{"name": "sample_sales.csv", "path": str(sample)}])

    def test_sample_total_is_564800(self):
        code = "df = dfs['sample_sales']\nresult = float(df['Sales'].sum())\nevidence = [{'rows': int(len(df))}]\nquality_notes = []"
        output = execute_plan(code, self.tables)
        self.assertEqual(output["result"], 564800.0)

    def test_profile_reports_shape(self):
        profile = profile_tables(self.tables, self.sources)["sample_sales"]
        self.assertEqual(profile["rows"], 12)
        self.assertIn("Sales", profile["columns"])

    def test_blocks_file_access(self):
        with self.assertRaises(ValueError):
            validate_code("result = open('secret')\nevidence = []\nquality_notes = []")

    def test_profit_is_refused_when_cost_data_is_missing(self):
        plan = create_plan(
            "What is the profit for each product?",
            profile_tables(self.tables, self.sources),
        )
        self.assertEqual(plan["status"], "CANNOT_DETERMINE")
        self.assertEqual(plan["code"], "")
        self.assertIn("no cost/profit data", plan["reason"])

    def test_uploaded_file_can_be_removed_from_workspace(self):
        with TemporaryDirectory() as media_root:
            with override_settings(MEDIA_ROOT=media_root):
                with patch(
                    "analyst.views.ollama_status",
                    return_value={"online": True, "ready": True, "model": "test"},
                ):
                    response = self.client.post(
                        reverse("upload"),
                        {
                            "datasets": SimpleUploadedFile(
                                "remove-me.csv",
                                b"region,sales\nNorth,10\n",
                                content_type="text/csv",
                            )
                        },
                    )
                    self.assertEqual(response.status_code, 302)
                    records = self.client.session["uploaded_files"]
                    self.assertEqual(len(records), 1)
                    stored_path = Path(records[0]["path"])
                    self.assertTrue(stored_path.is_file())

                    dashboard = self.client.get(reverse("dashboard"))
                    self.assertContains(dashboard, "Remove remove-me.csv")

                    response = self.client.post(reverse("delete_file", args=[0]))
                    self.assertEqual(response.status_code, 302)
                    self.assertFalse(stored_path.exists())
                    self.assertEqual(self.client.session["uploaded_files"], [])

    def test_upload_works_with_vercel_signed_cookie_sessions(self):
        with TemporaryDirectory() as media_root:
            with override_settings(
                MEDIA_ROOT=media_root,
                SESSION_ENGINE="django.contrib.sessions.backends.signed_cookies",
            ):
                response = self.client.post(
                    reverse("upload"),
                    {
                        "datasets": SimpleUploadedFile(
                            "vercel.csv",
                            b"region,sales\nNorth,10\n",
                            content_type="text/csv",
                        )
                    },
                )

                self.assertEqual(response.status_code, 302)
                session = self.client.session
                self.assertRegex(session["workspace_id"], r"^[0-9a-f]{32}$")
                self.assertEqual(len(session["uploaded_files"]), 1)
                self.assertTrue(Path(session["uploaded_files"][0]["path"]).is_file())

    def test_xlsx_can_be_removed_after_dashboard_profiles_it(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["region", "sales"])
        sheet.append(["North", 10])
        payload = BytesIO()
        workbook.save(payload)
        workbook.close()

        with TemporaryDirectory() as media_root:
            with override_settings(MEDIA_ROOT=media_root):
                with patch(
                    "analyst.views.ollama_status",
                    return_value={"online": True, "ready": True, "model": "test"},
                ):
                    response = self.client.post(
                        reverse("upload"),
                        {
                            "datasets": SimpleUploadedFile(
                                "remove-me.xlsx",
                                payload.getvalue(),
                                content_type=(
                                    "application/vnd.openxmlformats-officedocument."
                                    "spreadsheetml.sheet"
                                ),
                            )
                        },
                    )
                    self.assertEqual(response.status_code, 302)
                    stored_path = Path(
                        self.client.session["uploaded_files"][0]["path"]
                    )

                    dashboard = self.client.get(reverse("dashboard"))
                    self.assertContains(dashboard, "Remove remove-me.xlsx")

                    response = self.client.post(reverse("delete_file", args=[0]))

                    self.assertEqual(response.status_code, 302)
                    self.assertFalse(stored_path.exists())
                    self.assertEqual(self.client.session["uploaded_files"], [])

    def test_delete_rejects_files_outside_current_session_upload_folder(self):
        with TemporaryDirectory() as media_root, TemporaryDirectory() as other_root:
            with override_settings(MEDIA_ROOT=media_root):
                session = self.client.session
                session["uploaded_files"] = [
                    {
                        "name": "private.csv",
                        "path": str(Path(other_root) / "private.csv"),
                        "size": 1,
                    }
                ]
                session.save()
                protected_file = Path(other_root) / "private.csv"
                protected_file.write_bytes(b"protected")

                response = self.client.post(reverse("delete_file", args=[0]))

                self.assertEqual(response.status_code, 302)
                self.assertTrue(protected_file.is_file())
                self.assertEqual(len(self.client.session["uploaded_files"]), 1)

    def test_delete_allows_upload_from_previous_session_folder(self):
        with TemporaryDirectory() as media_root:
            with override_settings(MEDIA_ROOT=media_root):
                session = self.client.session
                session["uploaded_files"] = []
                session.save()
                old_folder = Path(media_root) / "uploads" / "old-session-key"
                old_folder.mkdir(parents=True)
                old_path = old_folder / "previous-upload.csv"
                old_path.write_bytes(b"region,sales\nNorth,10\n")
                session = self.client.session
                session["uploaded_files"] = [
                    {
                        "name": "previous-upload.csv",
                        "path": str(old_path),
                        "size": old_path.stat().st_size,
                    }
                ]
                session.save()

                response = self.client.post(reverse("delete_file", args=[0]))

                self.assertEqual(response.status_code, 302)
                self.assertFalse(old_path.exists())
                self.assertEqual(self.client.session["uploaded_files"], [])
